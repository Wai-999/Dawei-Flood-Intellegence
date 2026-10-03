"""Read the live Sheet export without changing or repairing source cells."""
import collections
import hashlib
import json
from pathlib import Path
import sys
import os
import openpyxl

path = Path(sys.argv[1])
out = Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)
book = openpyxl.load_workbook(path, data_only=True)
raw = {}
for sheet in book:
    raw[sheet.title] = [
        {"row": row, "cells": {openpyxl.utils.get_column_letter(col): value
         for col, value in enumerate(values, 1) if value is not None}}
        for row, values in enumerate(sheet.values, 1) if any(v is not None for v in values)
    ]
field = raw["ကွင်းဆင်းမှတ်တမ်းများ"][1:]
public = raw["ဘေးသင့်ဒေသစာရင်း"][1:]
named = [r for r in public if r["cells"].get("F") != "အချက်အလက်ဖြည့်စွက်ရန်"]
numeric_fields = list("FGHIJKLMNOP")
profile = {
    "source_url": os.environ.get("FLOOD_SOURCE_URL", "Private source snapshot; URL not supplied"),
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "sheets": {s.title: {"populated_rows": len(raw[s.title]), "max_columns": s.max_column} for s in book},
    "field_inventory_rows": len(field),
    "field_unique_location_keys": len({(r["cells"].get("D"), r["cells"].get("E")) for r in field}),
    "field_rows_with_assessment_data": sum(any(k not in {"A", "D", "E"} for k in r["cells"]) for r in field),
    "public_rows": len(public),
    "public_named_rows": len(named),
    "public_placeholder_rows": len(public)-len(named),
    "public_by_township": dict(collections.Counter(r["cells"].get("D") for r in named)),
    "public_numeric_block_nonempty_rows": sum(any(k in r["cells"] for k in "HIJKLMNO") for r in named),
    "public_nonnumeric_damage_rows": [r["row"] for r in named if "K" in r["cells"] and not isinstance(r["cells"]["K"], (float,int))],
    "public_numeric_verification_rows": [r["row"] for r in named if isinstance(r["cells"].get("O"),(float,int))],
    "coordinates_available": False,
    "complete_settlement_inventory_available": False,
    "summary_named_count_cell": "အနှစ်ချုပ်!B4",
    "summary_named_count": book["အနှစ်ချုပ်"]["B4"].value,
    "policy": "Public-report cells H:O retained for review, excluded from typed casualty/damage/need metrics due column alignment risk. No source repair or automatic verification."
}
(out / "source_snapshot.json").write_text(json.dumps(raw, ensure_ascii=False, indent=2, default=str))
(out / "source_profile.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2, default=str))
print(json.dumps(profile, ensure_ascii=False, indent=2, default=str))
